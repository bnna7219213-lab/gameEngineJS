# -*- coding: utf-8 -*-
import glob, io, sys

from pptx import Presentation

src = r"build\qa.pptx"
prs = Presentation(src)
out = io.StringIO()
for i, slide in enumerate(prs.slides, 1):
    out.write("=== SLIDE %d ===\n" % i)
    for shape in slide.shapes:
        if shape.has_text_frame:
            t = shape.text_frame.text.strip()
            if t:
                out.write(t + "\n")
        if shape.has_table:
            for row in shape.table.rows:
                out.write(" | ".join(c.text for c in row.cells) + "\n")
with open(r"build\qa_pptx.txt", "w", encoding="utf-8") as f:
    f.write(out.getvalue())
print("slides:", len(prs.slides.__iter__.__self__._sldIdLst) if False else len(prs.slides._sldIdLst))
print("text chars:", len(out.getvalue()))
